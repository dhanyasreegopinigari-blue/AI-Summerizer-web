from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    send_file
)

from flask_mail import Mail, Message

import smtplib
from email.mime.text import MIMEText
from docx import Document
from pptx import Presentation
from deep_translator import GoogleTranslator
from flask_login import (
    LoginManager,
    login_user,
    logout_user,
    login_required,
    current_user
)

from sqlalchemy import text
from werkzeug.security import (
    generate_password_hash,
    check_password_hash
)

from transformers import pipeline
import yake
summarizer = None

from config import Config
from models import db, User, Summary

from PyPDF2 import PdfReader
from reportlab.pdfgen import canvas
import io

app = Flask(__name__)
app.config.from_object(Config)

mail = Mail(app)

db.init_app(app)



login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "login"


@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))


def get_summarizer():
    global summarizer

    if summarizer is None:
        summarizer = pipeline(
            "summarization",
            model="sshleifer/distilbart-cnn-12-6"
        )

    return summarizer


@app.route("/")
def landing():
    return render_template("landing.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":

        username = request.form["username"]
        email = request.form["email"]
        password = request.form["password"]

        if User.query.filter_by(email=email).first():
            flash("Email already exists")
            return redirect(url_for("register"))

        user = User(
            username=username,
            email=email,
            password=generate_password_hash(password)
        )

        db.session.add(user)
        db.session.commit()

        flash("Account Created Successfully")
        return redirect(url_for("login"))

    return render_template("register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":

        email = request.form["email"]
        password = request.form["password"]

        user = User.query.filter_by(email=email).first()

        if user and check_password_hash(user.password, password):
            login_user(user)
            return redirect(url_for("dashboard"))

        flash("Invalid Credentials")

    return render_template("login.html")


@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("landing"))


@app.route("/dashboard")
@login_required
def dashboard():

    summaries = Summary.query.filter_by(
        user_id=current_user.id
    ).all()

    return render_template(
        "dashboard.html",
        summaries=summaries,
        document_count=len(summaries),
        words_saved=sum(
            len(s.summary_text.split())
            for s in summaries
        ),
        compression_rate=0,
        flashcards=[],
        keywords=[],
        important_sentences=[],
        mindmap="",
        score=None,
        readability=None
    )


@app.route("/summarize", methods=["POST"])
@login_required
def summarize():

    # TEXT INPUT
    text = request.form.get("text", "").strip()

    uploaded_file = request.files.get("file")

    # FILE PROCESSING
    if uploaded_file and uploaded_file.filename:
        filename = uploaded_file.filename.lower()
        text = ""

        if filename.endswith(".pdf"):
            pdf = PdfReader(uploaded_file)
            for page in pdf.pages:
                if page.extract_text():
                    text += page.extract_text() + "\n"

        elif filename.endswith(".docx"):
            doc = Document(uploaded_file)
            for p in doc.paragraphs:
                text += p.text + "\n"

        elif filename.endswith(".txt"):
            text = uploaded_file.read().decode("utf-8", errors="ignore")

        elif filename.endswith(".pptx"):
            ppt = Presentation(uploaded_file)
            for slide in ppt.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"

    if not text:
        flash("Please enter text or upload a file")
        return redirect(url_for("dashboard"))

    # SUMMARY TYPE
    summary_type = request.form.get("summary_type", "executive")

    if summary_type == "executive":
        max_len, min_len = 100, 30
    elif summary_type == "bullet":
        max_len, min_len = 80, 20
    elif summary_type == "research":
        max_len, min_len = 220, 80
    elif summary_type == "student":
        max_len, min_len = 150, 50
    elif summary_type == "meeting":
        max_len, min_len = 120, 40
    elif summary_type == "action":
        max_len, min_len = 80, 20
    else:
        max_len, min_len = 120, 30

    text = text[:3000]
        
    try:
        summarizer = get_summarizer()

        result = summarizer(
            text,
            max_length=max_len,
            min_length=min_len,
            do_sample=False
        )

        summary_text = result[0]["summary_text"]

    except Exception as e:
            flash(f"Error: {str(e)}")
            return redirect(url_for("dashboard"))

    # TRANSLATION
    language = request.form.get("language", "en")
    if language != "en":
        summary_text = GoogleTranslator(
            source="auto",
            target=language
        ).translate(summary_text)

    # SENTENCES
    sentences = summary_text.split(".")

    flashcards = []
    important_sentences = []
    mindmap = ""

    for s in sentences[:5]:
        s = s.strip()
        if len(s) > 15:
            flashcards.append({
                "question": f"What is meant by: {s[:40]}?",
                "answer": s
            })

        if len(s) > 20:
            important_sentences.append(s)

        if s:
            mindmap += f"├── {s}\n"

    # KEYWORDS
    kw_extractor = yake.KeywordExtractor()
    keywords = [k for k, v in kw_extractor.extract_keywords(text)[:10]]

    # SAVE DB
    new_summary = Summary(
        original_text=text,
        summary_text=summary_text,
        user_id=current_user.id
    )

    db.session.add(new_summary)
    db.session.commit()

    summaries = Summary.query.filter_by(user_id=current_user.id).all()

    return render_template(
        "dashboard.html",
        summary=summary_text,
        summaries=summaries,
        document_count=len(summaries),
        words_saved=sum(len(s.summary_text.split()) for s in summaries),
        compression_rate=70,
        score=90,
        readability=85,
        keywords=keywords,
        important_sentences=important_sentences,
        flashcards=flashcards,
        mindmap=mindmap
    )

@app.route("/ask-document", methods=["POST"])
@login_required
def ask_document():

    question = request.form.get("question")

    flash(f"You asked: {question}")

    return redirect(url_for("dashboard"))

@app.route("/download-pdf")
@login_required
def download_pdf():
    latest = Summary.query.filter_by(user_id=current_user.id).order_by(Summary.id.desc()).first()

    if not latest:
        return redirect(url_for("dashboard"))

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)

    pdf.drawString(100, 800, "AI Generated Summary")
    pdf.drawString(100, 760, latest.summary_text[:1000])

    pdf.save()
    buffer.seek(0)

    return send_file(buffer, as_attachment=True,
                     download_name="summary.pdf",
                     mimetype="application/pdf")

@app.route("/download-audio")
@login_required
def download_audio():
    flash("Audio feature coming soon")
    return redirect(url_for("dashboard"))

@app.route("/download-docx")
@login_required
def download_docx():

    latest = Summary.query.filter_by(
        user_id=current_user.id
    ).order_by(
        Summary.id.desc()
    ).first()

    if not latest:
        return redirect(url_for("dashboard"))

    doc = Document()
    doc.add_heading("AI Generated Summary", level=1)
    doc.add_paragraph(latest.summary_text)

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    return send_file(
        file_stream,
        as_attachment=True,
        download_name="summary.docx",
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

@app.route("/download-md")
@login_required
def download_md():

    latest = Summary.query.filter_by(
        user_id=current_user.id
    ).order_by(
        Summary.id.desc()
    ).first()

    if not latest:
        return redirect(url_for("dashboard"))

    markdown = f"# AI Generated Summary\n\n{latest.summary_text}"

    return send_file(
        io.BytesIO(markdown.encode()),
        as_attachment=True,
        download_name="summary.md",
        mimetype="text/markdown"
    )

@app.route("/send-email")
@login_required
def send_email():

    latest = Summary.query.filter_by(
        user_id=current_user.id
    ).order_by(
        Summary.id.desc()
    ).first()

    if not latest:
        flash("No summary found")
        return redirect(url_for("dashboard"))

    try:
        msg = Message(
            subject="Your AI Generated Summary",
            sender=app.config["MAIL_USERNAME"],
            recipients=[current_user.email]
        )

        msg.body = latest.summary_text

        mail.send(msg)

        flash("Summary sent to your email successfully!")

    except Exception as e:
        flash(f"Email error: {str(e)}")

    return redirect(url_for("dashboard"))

@app.route("/delete/<int:id>", methods=["POST"])
@login_required
def delete(id):
    item = Summary.query.get_or_404(id)

    if item.user_id != current_user.id:
        return redirect(url_for("dashboard"))

    db.session.delete(item)
    db.session.commit()

    return redirect(url_for("dashboard"))
  
if __name__ == "__main__":
    with app.app_context():
        db.create_all()

import os

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 10000))
    app.run(host="0.0.0.0", port=port)